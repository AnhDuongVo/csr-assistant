"""Source material: the draft CSR (Markdown or Word), source tables (CSV) and the protocol synopsis.

Every table row gets an id like T3.R2. Drafted sentences cite rows; the reviewer checks numbers against
exactly those rows.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

HEADING_RE = re.compile(r"^(#{1,6})\s+(?:(\d+(?:\.\d+)*)\.?\s*(?=[A-Za-z(]))?(.+?)\s*$")


@dataclass
class DocSection:
    number: str | None
    title: str
    level: int
    text: str
    line: int

    @property
    def word_count(self) -> int:
        return len(re.findall(r"\w+", self.text))


def parse_markdown(md: str) -> list[DocSection]:
    sections: list[DocSection] = []
    current: DocSection | None = None
    buf: list[str] = []
    for i, line in enumerate(md.splitlines(), start=1):
        m = HEADING_RE.match(line)
        if m:
            if current is not None:
                current.text = "\n".join(buf).strip()
                sections.append(current)
            current = DocSection(m.group(2), m.group(3).strip(), len(m.group(1)), "", i)
            buf = []
        else:
            buf.append(line)
    if current is not None:
        current.text = "\n".join(buf).strip()
        sections.append(current)
    return sections


def docx_to_markdown(path: str | Path) -> str:
    """Word drafts (what medical writers actually use): Heading N styles become Markdown headings."""
    try:
        import docx
    except ImportError as err:  # pragma: no cover
        raise RuntimeError("pip install -e '.[docx]' to read Word files") from err
    out = []
    for p in docx.Document(str(path)).paragraphs:
        style = (p.style.name or "").lower()
        m = re.match(r"heading (\d)", style)
        if m:
            out.append("#" * int(m.group(1)) + " " + p.text.strip())
        elif p.text.strip():
            out.append(p.text.strip())
        out.append("")
    return "\n".join(out)


def load_draft(path: str | Path) -> str:
    path = Path(path)
    return docx_to_markdown(path) if path.suffix.lower() == ".docx" else path.read_text(encoding="utf-8")


@dataclass
class Table:
    id: str
    name: str  # file stem, e.g. t14_3_1_ae_overview
    title: str
    kind: str  # disposition, demographics, efficacy, exposure, adverse_events, other
    columns: list[str]
    rows: list[list[str]]

    def row_id(self, r: int) -> str:
        return f"{self.id}.R{r + 1}"

    def render(self) -> str:
        head = f"{self.id}: {self.title} [{self.name}]\n| row | " + " | ".join(self.columns) + " |"
        lines = [head, "|" + "---|" * (len(self.columns) + 1)]
        for r, row in enumerate(self.rows):
            lines.append(f"| {self.row_id(r)} | " + " | ".join(row) + " |")
        return "\n".join(lines)

    def row_text(self, rid: str) -> str | None:
        m = re.fullmatch(rf"{re.escape(self.id)}\.R(\d+)", rid)
        if not m or not (1 <= int(m.group(1)) <= len(self.rows)):
            return None
        row = self.rows[int(m.group(1)) - 1]
        return " | ".join(f"{c}: {v}" for c, v in zip(self.columns, row, strict=False))


KIND_HINTS = {  # checked in this order
    "disposition": ("disposition",),
    "efficacy": ("efficacy", "endpoint"),
    "demographics": ("demograph", "baseline characteristics"),
    "exposure": ("exposure",),
    "adverse_events": ("_ae", "adverse", "teae", "safety"),
}


def _kind(name: str, title: str) -> str:
    low = f"{name} {title}".lower()
    for kind, hints in KIND_HINTS.items():
        if any(h in low for h in hints):
            return kind
    return "other"


def load_tables(directory: str | Path) -> list[Table]:
    """Each CSV is one table. An optional first line `# title: ...` gives the table title."""
    tables = []
    for i, path in enumerate(sorted(Path(directory).glob("*.csv")), start=1):
        lines = path.read_text(encoding="utf-8").splitlines()
        title = path.stem.replace("_", " ")
        if lines and lines[0].startswith("#"):
            title = lines[0].lstrip("#").split(":", 1)[-1].strip()
            lines = lines[1:]
        reader = list(csv.reader(lines))
        if not reader:
            continue
        tables.append(Table(f"T{i}", path.stem, title, _kind(path.stem, title), reader[0], reader[1:]))
    return tables


@dataclass
class SourcePackage:
    draft_md: str
    sections: list[DocSection]
    tables: list[Table]
    synopsis: str = ""
    name: str = "study"
    rows: dict[str, str] = field(default_factory=dict)  # row id -> "title | col: value | ..."
    row_values: dict[str, str] = field(default_factory=dict)  # row id -> values only (for n / N checks)
    row_labels: dict[str, str] = field(default_factory=dict)  # row id -> first cell + table title

    def __post_init__(self):
        for t in self.tables:
            for r, row in enumerate(t.rows):
                rid = t.row_id(r)
                self.rows[rid] = f"{t.title} | " + (t.row_text(rid) or "")
                self.row_values[rid] = " | ".join(row[1:] if len(row) > 1 else row)
                self.row_labels[rid] = f"{row[0] if row else ''} {t.title}"

    def tables_of(self, kinds: tuple[str, ...]) -> list[Table]:
        return [t for t in self.tables if t.kind in kinds]


def load_package(draft: str | Path, tables_dir: str | Path, synopsis: str | Path | None = None) -> SourcePackage:
    md = load_draft(draft)
    syn = Path(synopsis).read_text(encoding="utf-8") if synopsis and Path(synopsis).exists() else ""
    return SourcePackage(md, parse_markdown(md), load_tables(tables_dir), syn, Path(draft).stem)
