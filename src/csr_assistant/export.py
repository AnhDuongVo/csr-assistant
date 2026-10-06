"""Minimal Markdown to Word export, because medical writers work in Word."""

from __future__ import annotations

import re
from pathlib import Path


def markdown_to_docx(md: str, path: str | Path) -> Path:
    try:
        import docx
    except ImportError as err:  # pragma: no cover
        raise RuntimeError("pip install -e '.[docx]'") from err
    doc = docx.Document()
    for block in md.split("\n"):
        line = block.rstrip()
        if not line:
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            run = p.add_run(line.lstrip("> ").replace("**", ""))
            run.bold = True
        else:
            doc.add_paragraph(line.replace("**", ""))
    doc.save(str(path))
    return Path(path)
